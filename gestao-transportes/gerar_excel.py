"""Gera o ficheiro Gestao_Transportes.xlsx (modelo de gestão de transportes).

Uso: python gerar_excel.py
Para mudar nomes de condutores / matrículas antes de gerar, edite as listas abaixo.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

CONDUTORES = ["Condutor 1", "Condutor 2", "Condutor 3", "Condutor 4"]
CAMIOES = ["Camião 1", "Camião 2", "Camião 3", "Camião 4"]
ESTADOS = ["Pago", "Por pagar", "Parcial", "Em atraso"]
IVA_PADRAO = 0.23

GERAL_ROWS = 1000      # linhas preparadas na folha GERAL
VIEW_ROWS = 300        # linhas preparadas em cada folha de consulta
LIST_ROWS = 100        # tamanho das listas (clientes, fornecedores, ...)

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=14, color="1F3864")
BASE = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
GREY = Font(name=FONT, size=9, color="808080", italic=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
CALC_FILL = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
EUR = '#,##0.00 €;-#,##0.00 €;"-"'


def total_row(ws, row, label_col, first, last, sum_cols):
    """Linha de TOTAL no fundo da tabela. SUBTOTAL(109) respeita os filtros aplicados."""
    top = Side(style="double", color="1F3864")
    fill = PatternFill("solid", fgColor="D9E1F2")
    c = ws[f"{label_col}{row}"]
    c.value = "TOTAL"
    for col in sum_cols:
        ws[f"{col}{row}"] = f"=SUBTOTAL(109,{col}{first}:{col}{last})"
        ws[f"{col}{row}"].number_format = '#,##0.00 €;-#,##0.00 €;"-"'
    end = max(ws.max_column, 2)
    for j in range(1, end + 1):
        cell = ws.cell(row=row, column=j)
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.fill = fill
        cell.border = Border(top=top, bottom=thin, left=thin, right=thin)
DATE = "DD/MM/YYYY"
PCT = "0%"

# Colunas da folha GERAL: (cabeçalho, largura, formato, tipo)
# tipo: "in" = escrever à mão, "calc" = fórmula automática
COLS = [
    ("Data do Serviço", 13, DATE, "in"),
    ("Condutor", 16, None, "in"),
    ("Camião / Matrícula", 16, None, "in"),
    ("N.º Reboque / Contentor", 18, None, "in"),
    ("Cliente", 22, None, "in"),
    ("Fornecedor", 22, None, "in"),
    ("Local de Carga", 20, None, "in"),
    ("Local de Descarga", 20, None, "in"),
    ("Preço / Valor", 13, EUR, "in"),
    ("Taxa IVA", 9, PCT, "in"),
    ("IVA", 12, EUR, "calc"),
    ("Valor com IVA", 14, EUR, "calc"),
    ("N.º Fatura", 13, None, "in"),
    ("Data da Fatura", 13, DATE, "in"),
    ("Dia / Prazo de Pagamento", 15, DATE, "in"),
    ("Estado do Pagamento", 15, None, "in"),
    ("Data de Pagamento", 14, DATE, "in"),
    ("Observações", 30, None, "in"),
]
NC = len(COLS)
C = {h: L(i + 1) for i, (h, *_rest) in enumerate(COLS)}  # cabeçalho -> letra


def header_row(ws, row, headers, start_col=1):
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + j, value=h)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_status_cf(ws, rng, estado_col, prazo_col, first_row):
    """Cores discretas consoante o estado do pagamento."""
    e = f"${estado_col}{first_row}"
    p = f"${prazo_col}{first_row}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'{e}="Em atraso"'], fill=PatternFill("solid", fgColor="F4CCCC"),
        font=Font(color="9C0006"), stopIfTrue=True))
    # Por pagar / Parcial com prazo já ultrapassado -> também vermelho
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND(OR({e}="Por pagar",{e}="Parcial"),ISNUMBER({p}),{p}<TODAY())'],
        fill=PatternFill("solid", fgColor="F4CCCC"), font=Font(color="9C0006"), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'{e}="Por pagar"'], fill=PatternFill("solid", fgColor="FFF2CC"), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'{e}="Parcial"'], fill=PatternFill("solid", fgColor="FCE4D6"), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'{e}="Pago"'], fill=PatternFill("solid", fgColor="E2EFDA"), stopIfTrue=True))


wb = Workbook()

# ---------------------------------------------------------------- GERAL
g = wb.active
g.title = "GERAL"
header_row(g, 1, [h for h, *_ in COLS])
g.row_dimensions[1].height = 32
for j, (h, w, fmt, kind) in enumerate(COLS, start=1):
    g.column_dimensions[L(j)].width = w
for r in range(2, GERAL_ROWS + 2):
    for j, (h, w, fmt, kind) in enumerate(COLS, start=1):
        c = g.cell(row=r, column=j)
        c.font, c.border = BASE, BORDER
        if fmt:
            c.number_format = fmt
        if kind == "calc":
            c.fill = CALC_FILL
    I, J, K = C["Preço / Valor"], C["Taxa IVA"], C["IVA"]
    g[f"{K}{r}"] = f'=IF({I}{r}="","",{I}{r}*{J}{r})'
    g[f"{C['Valor com IVA']}{r}"] = f'=IF({I}{r}="","",{I}{r}+N({K}{r}))'
g.freeze_panes = "B2"
total_row(g, GERAL_ROWS + 2, "A", 2, GERAL_ROWS + 1,
          [C["Preço / Valor"], C["IVA"], C["Valor com IVA"]])
g.auto_filter.ref = f"A1:{L(NC)}{GERAL_ROWS + 1}"
g[f"{J}1"].comment = Comment(
    "Escreva a taxa de IVA da linha (ex.: 23%). O IVA e o Valor com IVA "
    "são calculados automaticamente (Valor com IVA = Valor + IVA). "
    "Se a fatura tiver outro IVA, pode escrever o valor do IVA diretamente por cima da fórmula.", "Modelo")
g[f"{C['Dia / Prazo de Pagamento']}1"].comment = Comment(
    "Data limite de pagamento. Se o estado for 'Por pagar' ou 'Parcial' e esta data já tiver passado, "
    "a linha fica a vermelho.", "Modelo")
add_status_cf(g, f"A2:{L(NC)}{GERAL_ROWS + 1}", C["Estado do Pagamento"],
              C["Dia / Prazo de Pagamento"], 2)

# ---------------------------------------------------------------- LISTAS
ls = wb.create_sheet("CLIENTES-FORNECEDORES")
ls["A1"] = "LISTAS — escreva aqui os nomes uma vez; aparecem nas listas suspensas da folha GERAL"
ls["A1"].font = TITLE_FONT
lists = [("Condutores", CONDUTORES), ("Camiões / Matrículas", CAMIOES),
         ("Reboques / Contentores", []), ("Clientes", []), ("Fornecedores", []),
         ("Estados do Pagamento", ESTADOS)]
names = ["L_Condutores", "L_Camioes", "L_Reboques", "L_Clientes", "L_Fornecedores", "L_Estados"]
LIST_REF = {}
for j, ((title, vals), nm) in enumerate(zip(lists, names), start=1):
    col = L(j * 2 - 1)  # A, C, E, G, I, K (coluna vazia entre listas)
    header_row(ls, 3, [title], start_col=j * 2 - 1)
    ls.column_dimensions[col].width = 26
    ls.column_dimensions[L(j * 2)].width = 3
    n = len(vals) if nm == "L_Estados" else LIST_ROWS
    for r in range(4, 4 + n):
        c = ls[f"{col}{r}"]
        c.font, c.border = BASE, BORDER
        if nm != "L_Estados":
            c.fill = INPUT_FILL
    for i, v in enumerate(vals):
        ls[f"{col}{4 + i}"] = v
    ref = f"'CLIENTES-FORNECEDORES'!${col}$4:${col}${3 + n}"
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
    LIST_REF[nm] = (col, ref)
ls["A2"] = ("Nota: se mudar o nome de um condutor ou matrícula aqui, a folha individual correspondente "
            "passa a filtrar pelo novo nome (pode renomear o separador à mão).")
ls["A2"].font = GREY
ls.freeze_panes = "A4"

# Listas suspensas na GERAL (aceitam texto novo, só avisam)
for nm, head in [("L_Condutores", "Condutor"), ("L_Camioes", "Camião / Matrícula"),
                 ("L_Reboques", "N.º Reboque / Contentor"), ("L_Clientes", "Cliente"),
                 ("L_Fornecedores", "Fornecedor")]:
    dv = DataValidation(type="list", formula1=f"={nm}", allow_blank=True,
                        errorStyle="warning", error="Este nome não está na lista. Deseja continuar?",
                        errorTitle="Nome novo")
    dv.add(f"{C[head]}2:{C[head]}{GERAL_ROWS + 1}")
    g.add_data_validation(dv)
dv = DataValidation(type="list", formula1="=L_Estados", allow_blank=True,
                    error="Escolha: Pago, Por pagar, Parcial ou Em atraso.", errorTitle="Estado inválido")
dv.add(f"{C['Estado do Pagamento']}2:{C['Estado do Pagamento']}{GERAL_ROWS + 1}")
g.add_data_validation(dv)
dv = DataValidation(type="date", operator="greaterThan", formula1="1", allow_blank=True,
                    error="Escreva uma data (DD/MM/AAAA).", errorTitle="Data inválida")
for h in ["Data do Serviço", "Data da Fatura", "Dia / Prazo de Pagamento", "Data de Pagamento"]:
    dv.add(f"{C[h]}2:{C[h]}{GERAL_ROWS + 1}")
g.add_data_validation(dv)


# ---------------------------------------------------------------- folhas de consulta
def view_sheet(title, label, filter_head, value_formula, list_name=None):
    """Folha que puxa automaticamente da GERAL as linhas onde <filter_head> = valor em C3."""
    ws = wb.create_sheet(title)
    ws["B1"] = f"{label} — serviços puxados automaticamente da folha GERAL"
    ws["B1"].font = TITLE_FONT
    ws["B3"] = f"{label}:"
    ws["B3"].font = BOLD
    ws["C3"] = value_formula
    ws["C3"].font = Font(name=FONT, size=12, bold=True, color="1F3864")
    ws["C3"].fill = INPUT_FILL
    ws["C3"].border = BORDER
    ws.merge_cells("C3:D3")
    if list_name:
        dv = DataValidation(type="list", formula1=f"={list_name}", allow_blank=True,
                            showErrorMessage=False)
        dv.add("C3")
        ws.add_data_validation(dv)
    ws["F3"] = "Não escrever nesta folha — os dados vêm da GERAL."
    ws["F3"].font = GREY

    HR = 6  # linha do cabeçalho
    F0 = HR + 1
    last = HR + VIEW_ROWS
    fcol = C[filter_head]
    gref = f"GERAL!${fcol}$2:${fcol}${GERAL_ROWS + 1}"

    # Totais
    ws["B4"] = "N.º de serviços:"
    ws["B4"].font = BOLD
    ws["C4"] = f'=IF($C$3="",0,COUNTIF({gref},$C$3))'
    ws["E4"] = "Total s/ IVA:"
    ws["E4"].font = BOLD
    ws["F4"] = f"=SUM(J{F0}:J{last})"
    ws["G4"] = "Total c/ IVA:"
    ws["G4"].font = BOLD
    ws["H4"] = f'=SUM(M{F0}:M{last})'
    ws["I4"] = "Por receber:"
    ws["I4"].font = BOLD
    ws["J4"] = f'=SUMIFS(M{F0}:M{last},Q{F0}:Q{last},"<>Pago")'
    for a in ("F4", "H4", "J4"):
        ws[a].number_format = EUR
        ws[a].font = BOLD

    # Coluna A = n.º da linha na GERAL (auxiliar), colunas B.. = colunas da GERAL
    header_row(ws, HR, ["Linha GERAL"] + [h for h, *_ in COLS])
    ws.row_dimensions[HR].height = 32
    ws.column_dimensions["A"].width = 7
    for j, (h, w, fmt, kind) in enumerate(COLS, start=2):
        ws.column_dimensions[L(j)].width = w
    ws.column_dimensions["A"].hidden = False
    for r in range(F0, last + 1):
        a = ws.cell(row=r, column=1)
        if r == F0:
            a.value = f'=IF($C$3="","",IFERROR(MATCH($C$3,{gref},0),""))'
        else:
            # procura a próxima ocorrência a partir da linha anterior
            a.value = (f'=IF(A{r-1}="","",IFERROR(A{r-1}+MATCH($C$3,'
                       f'INDEX(GERAL!${fcol}$2:${fcol}${GERAL_ROWS + 1},A{r-1}+1):'
                       f'GERAL!${fcol}${GERAL_ROWS + 1},0),""))')
        a.font = GREY
        a.border = BORDER
        for j, (h, w, fmt, kind) in enumerate(COLS, start=2):
            gcol = L(j - 1)
            src = f"INDEX(GERAL!${gcol}$2:${gcol}${GERAL_ROWS + 1},$A{r})"
            c = ws.cell(row=r, column=j, value=f'=IF($A{r}="","",IF({src}="","",{src}))')
            c.font, c.border = BASE, BORDER
            if fmt:
                c.number_format = fmt
    ws.freeze_panes = f"C{F0}"
    total_row(ws, last + 1, "B", F0, last, ["J", "L", "M"])
    ws.auto_filter.ref = f"A{HR}:{L(NC + 1)}{last}"
    add_status_cf(ws, f"B{F0}:{L(NC + 1)}{last}", "Q", "P", F0)
    return ws


# Colunas na folha de consulta: deslocadas 1 para a direita (B = Data ... M = Valor com IVA, Q = Estado)
assert C["Valor com IVA"] == "L" and C["Estado do Pagamento"] == "P" and C["Preço / Valor"] == "I"



cond_col, _ = LIST_REF["L_Condutores"]
cam_col, _ = LIST_REF["L_Camioes"]
for i, nome in enumerate(CONDUTORES):
    ws = view_sheet(nome.upper(), "Condutor", "Condutor",
                    f"='CLIENTES-FORNECEDORES'!{cond_col}{4 + i}")
for i, mat in enumerate(CAMIOES):
    ws = view_sheet(mat.upper(), "Camião / Matrícula", "Camião / Matrícula",
                    f"='CLIENTES-FORNECEDORES'!{cam_col}{4 + i}")
ws = view_sheet("REBOQUES-CONTENTORES", "N.º Reboque / Contentor", "N.º Reboque / Contentor",
                None, list_name="L_Reboques")
ws["C3"].comment = Comment("Escolha da lista ou escreva o n.º do reboque/contentor a consultar.", "Modelo")
ws = view_sheet("POR CLIENTE", "Cliente", "Cliente", None, list_name="L_Clientes")
ws["C3"].comment = Comment("Escolha da lista ou escreva o nome do cliente a consultar.", "Modelo")

# ---------------------------------------------------------------- RESUMO
rs = wb.create_sheet("RESUMO", 1)
rs["A1"] = "RESUMO — calculado automaticamente a partir da folha GERAL"
rs["A1"].font = TITLE_FONT
G = lambda h: f"GERAL!${C[h]}$2:${C[h]}${GERAL_ROWS + 1}"
VAL, TOT, EST = G("Preço / Valor"), G("Valor com IVA"), G("Estado do Pagamento")
DATA = G("Data do Serviço")

rs["A3"] = "Totais gerais"
rs["A3"].font = BOLD
header_row(rs, 4, ["Indicador", "Valor"])
geral = [
    ("Total de serviços", f'=COUNTIF({DATA},">0")+COUNTIFS({DATA},"",{VAL},"<>")', "0"),
    ("Total faturado (sem IVA)", f"=SUM({VAL})", EUR),
    ("Total IVA", f"=SUM({G('IVA')})", EUR),
    ("Total com IVA", f"=SUM({TOT})", EUR),
    ("Total pago", f'=SUMIFS({TOT},{EST},"Pago")', EUR),
    ("Total por pagar", f'=SUMIFS({TOT},{EST},"Por pagar")', EUR),
    ("Total parcial (valor total das faturas)", f'=SUMIFS({TOT},{EST},"Parcial")', EUR),
    ("Total em atraso", f'=SUMIFS({TOT},{EST},"Em atraso")', EUR),
    ("Sem estado indicado", f'=SUMIFS({TOT},{EST},"")', EUR),
]
for i, (lab, f, fmt) in enumerate(geral, start=5):
    rs[f"A{i}"], rs[f"B{i}"] = lab, f
    rs[f"B{i}"].number_format = fmt
    for a in (f"A{i}", f"B{i}"):
        rs[a].font, rs[a].border = BASE, BORDER
rs["A5"].comment = Comment("Conta as linhas da GERAL com Data do Serviço preenchida "
                           "(mais as que têm preço mas ainda sem data).", "Modelo")
rs.column_dimensions["A"].width = 34
rs.column_dimensions["B"].width = 16


def resumo_block(top, col0, title, list_name, gh, n):
    lcol, _ = LIST_REF[list_name]
    rs.cell(row=top, column=col0, value=title).font = BOLD
    hdr = [title.replace("Total por ", "").capitalize(), "N.º serviços", "Sem IVA", "Com IVA", "Pago", "Por receber"]
    header_row(rs, top + 1, hdr, start_col=col0)
    key = G(gh)
    for k in range(n):
        r = top + 2 + k
        nm = f"{L(col0)}{r}"
        rs[nm] = f"=IF('CLIENTES-FORNECEDORES'!{lcol}{4 + k}=\"\",\"\",'CLIENTES-FORNECEDORES'!{lcol}{4 + k})"
        f = [
            f'=IF({nm}="","",COUNTIF({key},{nm}))',
            f'=IF({nm}="","",SUMIFS({VAL},{key},{nm}))',
            f'=IF({nm}="","",SUMIFS({TOT},{key},{nm}))',
            f'=IF({nm}="","",SUMIFS({TOT},{key},{nm},{EST},"Pago"))',
            f'=IF({nm}="","",SUMIFS({TOT},{key},{nm})-SUMIFS({TOT},{key},{nm},{EST},"Pago"))',
        ]
        for j, ff in enumerate(f, start=1):
            c = rs.cell(row=r, column=col0 + j, value=ff)
            c.number_format = "0" if j == 1 else EUR
        for j in range(6):
            c = rs.cell(row=r, column=col0 + j)
            c.font, c.border = BASE, BORDER
    return top + 2 + n + 1


nxt = resumo_block(16, 1, "Total por condutor", "L_Condutores", "Condutor", 10)
nxt = resumo_block(nxt, 1, "Total por camião", "L_Camioes", "Camião / Matrícula", 10)
nxt = resumo_block(nxt, 1, "Total por cliente", "L_Clientes", "Cliente", LIST_ROWS)
for col in "CDEF":
    rs.column_dimensions[col].width = 15
rs["D3"] = "Os nomes vêm da folha CLIENTES-FORNECEDORES. Linhas vazias = nomes ainda não preenchidos."
rs["D3"].font = GREY
rs.freeze_panes = "A2"

# ---------------------------------------------------------------- COMO USAR
hw = wb.create_sheet("COMO USAR")
txt = [
    ("COMO USAR ESTE FICHEIRO", TITLE_FONT),
    ("", BASE),
    ("1. Na folha CLIENTES-FORNECEDORES escreva os nomes dos condutores, matrículas, reboques/contentores, clientes e fornecedores.", BASE),
    ("   Os separadores CONDUTOR 1..4 e CAMIÃO 1..4 filtram automaticamente pelo 1.º..4.º nome dessas listas (pode renomear os separadores à mão).", BASE),
    ("2. Registe cada serviço UMA vez na folha GERAL (uma linha por serviço). Células a cinzento (IVA, Valor com IVA) são automáticas.", BASE),
    ("   Taxa IVA: escreva por ex. 23%. Valor com IVA = Valor + IVA. Se a fatura tiver um IVA diferente, escreva o valor do IVA por cima da fórmula.", BASE),
    ("3. As folhas dos condutores, camiões, REBOQUES-CONTENTORES e POR CLIENTE mostram automaticamente os serviços da GERAL (não escrever nelas).", BASE),
    ("   Em REBOQUES-CONTENTORES e POR CLIENTE escolha/escreva o valor na célula amarela (C3).", BASE),
    ("4. A folha RESUMO calcula totais gerais, por condutor, por camião e por cliente.", BASE),
    ("5. Estado do Pagamento: Pago (verde), Por pagar (amarelo), Parcial (laranja), Em atraso (vermelho).", BASE),
    ("   Uma linha 'Por pagar' ou 'Parcial' com o prazo de pagamento já ultrapassado também fica a vermelho.", BASE),
    ("6. Datas no formato DD/MM/AAAA. Valores em euros (€).", BASE),
    ("   No fundo de cada tabela há uma linha TOTAL (Valor, IVA e Valor com IVA). Atalho: Ctrl+End. Se filtrar, o total soma só as linhas visíveis.", BASE),
    ("7. Informação ilegível nas folhas em papel: escreva 'REVER' na coluna Observações em vez de adivinhar.", BASE),
    ("", BASE),
    ("Correspondência com as folhas em papel:", BOLD),
    ("   Folha de cima: Data → Data do Serviço · Fornecedor → Fornecedor · N.º Reboque → N.º Reboque / Contentor · Valor → Preço / Valor", BASE),
    ("   IVA → IVA · Com IVA → Valor com IVA · N.º Fatura → N.º Fatura · Data → Data da Fatura · Dia → Dia / Prazo de Pagamento · Pagamento → Estado do Pagamento / Data de Pagamento", BASE),
    ("   Folha de baixo: Data → Data do Serviço · Cliente → Cliente · Carga → Local de Carga · Descarga → Local de Descarga · Preço → Preço / Valor", BASE),
    ("   Se a mesma viagem aparece nas duas folhas, preencha tudo na mesma linha da GERAL.", BASE),
    ("", BASE),
    ("Exemplo de preenchimento (apenas ilustrativo — NÃO está na GERAL nem conta nos totais):", BOLD),
]
for i, (t, f) in enumerate(txt, start=1):
    hw[f"A{i}"] = t
    hw[f"A{i}"].font = f
ex_row = len(txt) + 2
exemplo = ["15/03/2026", "João Silva", "AA-00-BB", "L-123456", "Cliente Exemplo, Lda", "Fornecedor Exemplo",
           "Porto", "Lisboa", "450,00 €", "23%", "103,50 €", "553,50 €", "FT 2026/15", "20/03/2026",
           "19/04/2026", "Por pagar", "", "Exemplo"]
for j, (h, *_r) in enumerate(COLS):
    hw.cell(row=ex_row, column=2 + j, value=h)
    hw.cell(row=ex_row + 1, column=2 + j, value=exemplo[j]).font = GREY
header_row(hw, ex_row, [h for h, *_ in COLS], start_col=2)
for j, (h, w, *_r) in enumerate(COLS, start=2):
    hw.column_dimensions[L(j)].width = w
hw.column_dimensions["A"].width = 3
hw.sheet_view.showGridLines = False
for i in range(1, len(txt) + 1):
    hw[f"A{i}"].alignment = Alignment(wrap_text=False)

# Ordem final dos separadores
order = ["COMO USAR", "GERAL", "RESUMO"] + [n.upper() for n in CONDUTORES] + [m.upper() for m in CAMIOES] + \
        ["REBOQUES-CONTENTORES", "POR CLIENTE", "CLIENTES-FORNECEDORES"]
wb._sheets = [wb[n] for n in order]
wb.active = 1
g.sheet_properties.tabColor = "1F3864"
rs.sheet_properties.tabColor = "548235"
for n in CONDUTORES:
    wb[n.upper()].sheet_properties.tabColor = "8EA9DB"
for n in CAMIOES:
    wb[n.upper()].sheet_properties.tabColor = "A9D08E"

import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Gestao_Transportes.xlsx")
wb.save(out)
print("Gerado:", out)
